#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# STATUS: ATIVO-P7 (espelho principal; Onda 1.7-b JSON-canônico)
"""Gera a planilha consolidada da Base Terminológica RC (base-terminologica.xlsx)
a partir do canônico base/canonico.json (exportado de ferramentas/dados_termos.py).
Paridade 1.7-b: por CONTEÚDO (célula a célula) — o artefato é zip com timestamps.
Execute a partir de rc-terminologia/ :  python3 ferramentas/build_xlsx.py
"""
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__)))
from loader import carregar_canonico, BIBLIO  # Onda 1.7-b: JSON-primeiro
TERMOS, RELACOES = carregar_canonico()

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "base-terminologica.xlsx")

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=16, color="1F3864")
SUB_FONT = Font(italic=True, size=10, color="595959")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
STATUS_FILL = {"candidato": None, "provisório": "FFF2CC", "em análise": "FCE4D6", "verificado": "E2EFDA"}

STATUS_LABEL = {"candidato": "⚪ candidato", "provisório": "🟡 provisório",
                "em análise": "🟠 em análise", "verificado": "🟢 verificado"}

CATS = [
    ("Seres & Entidades", ["1.1 Divindades / Criadores", "1.2 Coletivos criadores", "1.3 Extraterrestres / Raças cósmicas",
                           "1.4 Anjos / Arcanjos / Querubins", "1.5 Decaídos / Exilados / Demônios",
                           "1.6 Figuras humanas (históricas/míticas)", "1.7 Híbridos / Seres de transição", "1.8 Consciências coletivas / Logos"]),
    ("Conceitos Cosmológicos", ["2.1 Cosmogonia / Criação", "2.2 Leis e princípios cósmicos", "2.3 Dimensões / Planos de existência",
                                "2.4 Espiritualidade / Doutrina", "2.5 Propósitos / Destinos", "2.6 Ética / Carma / Compromissos"]),
    ("Lokas & Geografias", ["3.1 Planetas / Sistemas estelares", "3.2 Lokas / Mundos interiores", "3.3 Localidades terrestres",
                            "3.4 Regiões cósmicas / Portais"]),
    ("Processos & Fenômenos", ["4.1 Processos biológicos / genéticos", "4.2 Fenômenos espirituais / mediúnicos",
                               "4.3 Transmigração / Reencarnação", "4.4 Fenômenos físicos / energéticos"]),
    ("Tecnologias & Artefatos", ["5.1 Naves / Tecnologia extraterrestre", "5.2 Artefatos sagrados / Relíquias",
                                 "5.3 Dispositivos / Instrumentos", "5.4 Escrituras / Objetos documentais"]),
    ("Genealogias & Linhagens", ["6.1 Famílias divinas / Celestiais", "6.2 Linhagens humanas", "6.3 Dinastias extraterrestres",
                                 "6.4 Árvores genealógicas"]),
    ("Eventos & Eras", ["7.1 Guerras / Rebeliões / Conflitos", "7.2 Eras / Épocas / Ciclos", "7.3 Profecias / Anúncios",
                        "7.4 Catástrofes / Dilúvios / Quedas", "7.5 Missões / Expedições / Encontros"]),
    ("Termos Específicos", ["8.1 Sânscrito / Hinduísmo", "8.2 Hebraico / Cabala", "8.3 Sumério / Acádio / Babilônio",
                            "8.4 Grego / Latim", "8.5 Outros idiomas / jargão técnico"]),
]



def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def finish(ws, header_row, widths):
    set_widths(ws, widths)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{max(header_row + 1, ws.max_row)}"


def write_rows(ws, start_row, rows, ncols, status_col=None):
    for r, row in enumerate(rows, start=start_row):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = WRAP
            cell.border = BORDER
            if status_col and c == status_col and val in STATUS_FILL and STATUS_FILL[val]:
                cell.fill = PatternFill("solid", fgColor=STATUS_FILL[val])


wb = Workbook()

# ------------------------------------------------------------------ Como usar
ws = wb.active
ws.title = "Como usar"
ws.sheet_view.showGridLines = False
ws["B2"] = "Base Terminológica — Revelações Cósmicas (Jan Val Ellam)"
ws["B2"].font = TITLE_FONT
ws["B4"] = ("Planilha consolidada e navegável do projeto. As fichas-fonte vivem em Markdown (pasta termos/ + indice-mestre.md); "
            "esta planilha é uma visão filtrada/consolidada e pode ser regenerada pelo script ferramentas/build_xlsx.py.")
ws["B4"].alignment = WRAP
ws["B4"].font = SUB_FONT
ws.column_dimensions["A"].width = 2
ws.column_dimensions["B"].width = 110

guias = [
    ("1", "Índice Mestre", "Todos os termos (1 linha cada), com status e filtros. É a porta de entrada."),
    ("2", "Fichas", "Ficha completa dos termos centrais (campos da ficha v2)."),
    ("3", "Categorias", "Vocabulário controlado: categorias e subcategorias (códigos 1.1, 7.3 etc.)."),
    ("4", "Bibliografia", "Cadastro das obras do corpus com siglas B###."),
    ("5", "Relações", "Grafo: Termo A — tipo de relação — Termo B (mini-ontologia)."),
    ("6", "Status", "Significado de cada status (candidato, provisório, em análise, verificado)."),
]
r = 6
ws.cell(row=r, column=2, value="Guias desta planilha:").font = Font(bold=True, size=12)
r += 1
for num, nome, desc in guias:
    ws.cell(row=r, column=2, value=f"[{num}] {nome} — {desc}").font = SUB_FONT
    r += 1

# ------------------------------------------------------------------ Índice Mestre
ws = wb.create_sheet("Índice Mestre")
headers = ["Código", "Termo", "Categoria", "Subcategoria", "Status", "Fontes (siglas)", "Termos relacionados", "Ficha completa?", "Última atualização"]
ws.append(headers)
style_header(ws, 1, len(headers))

indice = []
for cod, nome, cat, sub, status, fontes, rel, ficha in TERMOS:
    indice.append([cod, nome, cat, sub, STATUS_LABEL[status], ", ".join(fontes) if fontes else "—",
                   rel, "sim" if ficha else "não", "2026-08-14"])
write_rows(ws, 2, indice, len(headers), status_col=5)
finish(ws, 1, [10, 42, 24, 34, 16, 14, 30, 13, 16])

dv_status = DataValidation(type="list", formula1='"⚪ candidato,🟡 provisório,🟠 em análise,🟢 verificado"', allow_blank=True)
ws.add_data_validation(dv_status)
dv_status.add("E2:E600")

# ------------------------------------------------------------------ Fichas (resumo)
ws = wb.create_sheet("Fichas")
headers = ["Código", "Termo", "Categoria", "Subcategoria", "Status", "Definição sintética",
           "Contexto / Origem", "Etimologia / Grafias", "Fontes (obra, capítulo, página)",
           "Citações-chave", "Termos relacionados", "Seres associados", "Observações", "Última atualização"]
ws.append(headers)
style_header(ws, 1, len(headers))
fichas = []
for cod, nome, cat, sub, status, fontes, rel, ficha in TERMOS:
    if not ficha:
        continue
    fichas.append([cod, nome, cat, sub, STATUS_LABEL[status],
                   "(ver ficha completa em termos/)", "(ver ficha completa em termos/)",
                   "(ver ficha completa em termos/)", ", ".join(fontes), "(ver ficha completa em termos/)",
                   rel, "—", "(ver ficha completa em termos/)", "2026-08-14"])
write_rows(ws, 2, fichas, len(headers), status_col=5)
finish(ws, 1, [10, 34, 24, 30, 16, 30, 30, 30, 20, 30, 30, 20, 30, 15])

dv_cat = DataValidation(type="list",
                        formula1='"Seres & Entidades,Conceitos Cosmológicos,Lokas & Geografias,Processos & Fenômenos,Tecnologias & Artefatos,Genealogias & Linhagens,Eventos & Eras,Termos Específicos"',
                        allow_blank=True)
ws.add_data_validation(dv_cat)
dv_cat.add("C2:C600")

# ------------------------------------------------------------------ Categorias
ws = wb.create_sheet("Categorias")
headers = ["Código", "Categoria", "Subcategoria"]
ws.append(headers)
style_header(ws, 1, len(headers))
cat_rows = []
for cat, subs in CATS:
    for sub in subs:
        cat_rows.append([sub.split(" ")[0], cat, sub])
write_rows(ws, 2, cat_rows, len(headers))
finish(ws, 1, [10, 26, 44])

# ------------------------------------------------------------------ Bibliografia
ws = wb.create_sheet("Bibliografia")
headers = ["Código", "Título", "Subtítulo", "Editora", "Ano", "Status", "Observações"]
ws.append(headers)
style_header(ws, 1, len(headers))
write_rows(ws, 2, BIBLIO, len(headers))
finish(ws, 1, [10, 40, 44, 22, 10, 16, 70])

# ------------------------------------------------------------------ Relações
ws = wb.create_sheet("Relações")
headers = ["Termo A (código)", "Termo A (nome)", "Tipo de relação", "Termo B (código)", "Termo B (nome)", "Obra-fonte"]
ws.append(headers)
style_header(ws, 1, len(headers))
nome = {t[0]: t[1] for t in TERMOS}
rel_rows = []
for a, tipo, b, obra in RELACOES:
    rel_rows.append([a, nome[a], tipo, b, nome[b], obra])
write_rows(ws, 2, rel_rows, len(headers))
finish(ws, 1, [16, 34, 18, 16, 40, 12])

# ------------------------------------------------------------------ Status
ws = wb.create_sheet("Status")
headers = ["Status", "Significado", "Quando aplicar"]
ws.append(headers)
style_header(ws, 1, len(headers))
status_rows = [
    ["⚪ candidato", "Termo identificado, sem ficha ainda", "Só aparece no Índice Mestre"],
    ["🟡 provisório", "Definição documentada em fonte publicada, mas não verificada em obras primárias de Ellam", "Após extração de glossário/PDF"],
    ["🟠 em análise", "Ficha redigida, aguardando revisão e referências página a página", "Após cruzamento com obras primárias"],
    ["🟢 verificado", "Revisado e referenciado (obra + capítulo/página)", "Pode ser citado como consolidado"],
]
write_rows(ws, 2, status_rows, len(headers))
finish(ws, 1, [16, 70, 55])

wb.save(OUT)
print(f"OK -> {OUT} ({len(wb.sheetnames)} abas)")
print(f"Índice: {len(indice)} termos | Fichas-resumo: {len(fichas)} | Relações: {len(rel_rows)} | Obras: {len(BIBLIO)}")
