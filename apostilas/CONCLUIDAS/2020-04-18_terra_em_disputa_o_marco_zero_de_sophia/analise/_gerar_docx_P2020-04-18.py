#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o DOCX da apostila P2020-04-18 a partir do apostila.md.

Padrão do rito: @metadados no topo; # headings; blockquotes com estilo
Intense Quote; negrito **...**; lista de dicionário no Apêndice.
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

AQUI = Path(__file__).resolve().parent
MD = AQUI.parent / "apostila.md"
OUT = AQUI.parent / "Apostila - Terra em Disputa — o Marco Zero de Sophia.docx"

doc = Document()

# estilos base
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)

for sec in doc.sections:
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(3)
    sec.right_margin = Cm(2.5)


def add_runs(par, text):
    """Aplica **negrito** e *itálico* (simples) em runs."""
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*", text):
        antes = text[pos:m.start()]
        if antes:
            _add_italic_runs(par, antes)
        r = par.add_run(m.group(1))
        r.bold = True
        pos = m.end()
    resto = text[pos:]
    if resto:
        _add_italic_runs(par, resto)


def _add_italic_runs(par, text):
    pos = 0
    for m in re.finditer(r"(?<!\*)\*([^*\n]+?)\*(?!\*)", text):
        antes = text[pos:m.start()]
        if antes:
            par.add_run(antes)
        r = par.add_run(m.group(1))
        r.italic = True
        pos = m.end()
    resto = text[pos:]
    if resto:
        par.add_run(resto)


linhas = MD.read_text(encoding="utf-8").split("\n")
i = 0
while i < len(linhas):
    l = linhas[i].rstrip()
    if not l.strip():
        i += 1
        continue
    if l.startswith("@"):
        m = re.match(r"@(\w+):\s*(.*)", l)
        if m:
            p = doc.add_paragraph()
            r = p.add_run(f"{m.group(1).capitalize()}: ")
            r.bold = True
            add_runs(p, m.group(2))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if m.group(1) == "titulo" else WD_ALIGN_PARAGRAPH.LEFT
        i += 1
        continue
    if l.startswith("# "):
        doc.add_heading(l[2:].strip(), level=1)
        i += 1
        continue
    if l.startswith("## "):
        doc.add_heading(l[3:].strip(), level=2)
        i += 1
        continue
    if l.startswith("> "):
        p = doc.add_paragraph(style="Intense Quote")
        add_runs(p, l[2:].strip())
        i += 1
        continue
    if l.strip() == "---":
        doc.add_paragraph()
        i += 1
        continue
    # parágrafo comum
    p = doc.add_paragraph()
    add_runs(p, l.strip())
    i += 1

doc.save(OUT)
print("OK ->", OUT)
